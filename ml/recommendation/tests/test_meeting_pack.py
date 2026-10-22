"""Tests for the NEBSL meeting workbook. Built from the config only (no data)."""
import pytest
from openpyxl import load_workbook

from recommendation.config import load_config
from recommendation.matching.ahp import PAIRS, matrix, read_form, weights_and_cr
from recommendation.reporting.meeting_pack import SUS, build_pack


@pytest.fixture()
def pack(tmp_path):
    out = tmp_path / "pack.xlsx"
    counts = build_pack(load_config(), out)
    return out, counts


def test_pack_has_all_sheets(pack):
    out, counts = pack
    assert list(counts) == ["0_Read_me", "1_AHP", "2_Thresholds", "3_Usability"]
    assert load_workbook(out).sheetnames == list(counts)
    assert counts["1_AHP"] == len(PAIRS)


def test_filled_ahp_sheet_is_read_back_by_the_ahp_command(pack):
    out, _ = pack
    wb = load_workbook(out)
    ws = wb["1_AHP"]
    for i in range(len(PAIRS)):
        ws.cell(4 + i, 4).value, ws.cell(4 + i, 5).value = "Equal", 1          # clinician 1
        ws.cell(4 + i, 6).value, ws.cell(4 + i, 7).value = "A", 3              # clinician 2
    wb.save(out)
    sets = read_form(out)
    assert len(sets) == 2
    w, cr = weights_and_cr(matrix(sets[0]))
    assert all(v == pytest.approx(0.2) for v in w.values()) and cr == pytest.approx(0, abs=1e-6)


def test_empty_ahp_sheet_is_reported(pack):
    out, _ = pack
    with pytest.raises(ValueError, match="No answers"):
        read_form(out)


def test_thresholds_come_from_the_config(pack):
    out, _ = pack
    cfg = load_config()["followup"]["alerts"]
    values = {(r[0], r[1]): r[2] for r in load_workbook(out)["2_Thresholds"].iter_rows(min_row=3, values_only=True)}
    assert values[("Red alert", "failure risk 12m")] == cfg["red"]["failure_risk_12m"]
    assert values[("Amber alert", "loss 12m at least")] == cfg["amber"]["loss_12m_at_least"]
    assert values[("Next visit", "red alert")] == cfg["interval_months"]["red"]


def test_usability_sheet_has_sus_and_score_formula(pack):
    out, _ = pack
    rows = list(load_workbook(out)["3_Usability"].iter_rows(values_only=True))
    statements = [r[1] for r in rows if r[1] in SUS]
    assert statements == SUS
    score = next(r for r in rows if r[1] == "SUS score (0-100)")
    assert str(score[2]).startswith("=IF(COUNT(") and "*2.5" in score[2]
