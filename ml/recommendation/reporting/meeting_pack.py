"""Workbook for the NEBSL clinician meeting (`python -m recommendation.main pack`).

Sheets:
  0_Read_me           what the meeting is for and how to fill the workbook
  1_AHP               pairwise comparisons of the five MCDA criteria, up to three clinicians;
                      read back by `main ahp` (matching/ahp.read_form) to compute the weights
  2_Thresholds        follow-up alert thresholds and data-warning limits (placeholders) to agree
  3_Usability         tasks to try in the app, and the System Usability Scale (SUS) per clinician

Built from config.yaml only: it holds no patient or tissue data.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from recommendation.matching.ahp import LABEL, PAIRS

CLINICIANS = 3
HEADER = Font(bold=True)
TITLE = Font(bold=True, size=13)
FILL = PatternFill("solid", fgColor="FFF4CC")     # cells the clinicians fill in
WRAP = Alignment(wrap_text=True, vertical="top")

# Data-warning limits used by the app (ml-service/pipelines/recommendation/checks.py and the request limits).
DATA_LIMITS = [
    ("Matching", "Usual storage limit (warning above)", 14, "days"),
    ("Matching", "Minimum donor ECD for PK / DSEK / DMEK (warning below)", 2000, "cells/mm²"),
    ("Matching", "Longest death-to-preservation time accepted", 72, "hours"),
    ("Matching", "Longest storage accepted", 60, "days"),
    ("Follow-up", "ECD rise between visits treated as a likely typing error (warning above)", 0.10, "fraction"),
    ("Follow-up", "Donor ECD below the usual eye-bank minimum (warning below)", 2000, "cells/mm²"),
]
TASKS = [
    "Rank the corneas in stock for a recipient request and say which one you would choose",
    "Explain from the screen why the #1 cornea is ranked above #2",
    "Find a cornea whose serology is not confirmed and say what would change its position",
    "Enter a graft's 12-month visit and read the alert, next visit and review prompts",
    "Explain from the screen what drives the graft's predicted risk",
    "Reopen an earlier ranking or follow-up from 'Previous runs'",
]
SUS = [  # Brooke (1996), with "system" as in the original wording
    "I think that I would like to use this system frequently.",
    "I found the system unnecessarily complex.",
    "I thought the system was easy to use.",
    "I think that I would need the support of a technical person to be able to use this system.",
    "I found the various functions in this system were well integrated.",
    "I thought there was too much inconsistency in this system.",
    "I would imagine that most people would learn to use this system very quickly.",
    "I found the system very cumbersome to use.",
    "I felt very confident using the system.",
    "I needed to learn a lot of things before I could get going with this system.",
]


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w


def _read_me(ws, cfg: dict) -> int:
    rows = [
        ("NEBSL clinician meeting: recommendation component",),
        ("",),
        ("Purpose", "Agree the clinical settings of the donor-cornea matching and post-transplant follow-up tools, "
                    "and record how usable the tools are."),
        ("1_AHP", "Each clinician compares every pair of matching criteria: which matters more, and how much "
                  "(1 = equal, 3 = moderately, 5 = strongly, 7 = very strongly, 9 = extremely). "
                  "Fill only the yellow cells. The weights are then computed with: main ahp --apply"),
        ("2_Thresholds", "For each alert threshold and data-warning limit, keep the current value or write the "
                         "value you propose, with a short reason."),
        ("3_Usability", "Try each task in the app, note whether it was completed and any comment, then answer the "
                        "ten SUS statements (1 = strongly disagree, 5 = strongly agree)."),
        ("Current MCDA weights", f"{cfg['matching']['weights_version']}: " + ", ".join(
            f"{k.replace('_', ' ')} {v:.2f}" for k, v in cfg["matching"]["weights"].items())),
        ("Alert rules version", str(cfg["followup"]["alerts"]["version"])),
        ("Data", "This workbook contains no patient or tissue data."),
    ]
    for r in rows:
        ws.append(list(r))
    ws["A1"].font = TITLE
    for row in ws.iter_rows(min_row=3):
        row[0].font = HEADER
        for cell in row:
            cell.alignment = WRAP
    _widths(ws, [24, 110])
    return len(rows) - 2


def _ahp(ws) -> int:
    ws.append(["Matching criteria: pairwise comparisons (AHP)"])
    ws.append(["For each pair: who is more important (A, B or Equal) and how much (1-9). Yellow cells only."])
    header = ["Pair", "Criterion A", "Criterion B"]
    for k in range(1, CLINICIANS + 1):
        header += [f"Clinician {k}: more important", f"Clinician {k}: how much (1-9)"]
    ws.append(header)   # row 3: read_form uses header=2
    winner = DataValidation(type="list", formula1='"A,B,Equal"', allow_blank=True)
    strength = DataValidation(type="whole", operator="between", formula1="1", formula2="9", allow_blank=True)
    ws.add_data_validation(winner)
    ws.add_data_validation(strength)
    for i, (a, b) in enumerate(PAIRS, start=1):
        ws.append([i, LABEL[a], LABEL[b]] + [None] * (2 * CLINICIANS))
        r = ws.max_row
        for k in range(CLINICIANS):
            win, how = ws.cell(r, 4 + 2 * k), ws.cell(r, 5 + 2 * k)
            win.fill = how.fill = FILL
            winner.add(win)
            strength.add(how)
    ws["A1"].font = TITLE
    for cell in ws[3]:
        cell.font = HEADER
        cell.alignment = WRAP
    _widths(ws, [6, 34, 34] + [16, 14] * CLINICIANS)
    return len(PAIRS)


def _thresholds(ws, cfg: dict) -> int:
    alerts = cfg["followup"]["alerts"]
    ws.append(["Follow-up alert thresholds and data-warning limits (placeholders until agreed)"])
    ws.append(["Area", "Setting", "Current value", "Unit", "Proposed value", "Reason"])
    rows = [("Red alert", k.replace("_", " "), v, "") for k, v in alerts["red"].items()]
    rows.append(("Red alert", "rejection episode by 12 months", "always red", "rule"))
    rows += [("Amber alert", k.replace("_", " "), v, "") for k, v in alerts["amber"].items()]
    rows += [("Next visit", f"{level} alert", months, "months") for level, months in alerts["interval_months"].items()]
    rows += DATA_LIMITS
    for area, setting, value, unit in rows:
        ws.append([area, setting, value, unit, None, None])
        ws.cell(ws.max_row, 5).fill = ws.cell(ws.max_row, 6).fill = FILL
    ws["A1"].font = TITLE
    for cell in ws[2]:
        cell.font = HEADER
    _widths(ws, [14, 62, 14, 12, 16, 40])
    return len(rows)


def _usability(ws) -> int:
    ws.append(["Usability: tasks in the app"])
    header = ["#", "Task"]
    for k in range(1, CLINICIANS + 1):
        header += [f"Clinician {k}: completed (Yes/No)", f"Clinician {k}: comment"]
    ws.append(header)
    done = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(done)
    for i, task in enumerate(TASKS, start=1):
        ws.append([i, task] + [None] * (2 * CLINICIANS))
        for k in range(CLINICIANS):
            ws.cell(ws.max_row, 3 + 2 * k).fill = ws.cell(ws.max_row, 4 + 2 * k).fill = FILL
            done.add(ws.cell(ws.max_row, 3 + 2 * k))
    tasks_header = 2

    ws.append([])
    ws.append(["System Usability Scale (SUS): 1 = strongly disagree ... 5 = strongly agree"])
    sus_title = ws.max_row
    ws.append(["#", "Statement"] + [f"Clinician {k}" for k in range(1, CLINICIANS + 1)])
    sus_header = ws.max_row
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(score)
    first = ws.max_row + 1
    for i, statement in enumerate(SUS, start=1):
        ws.append([i, statement] + [None] * CLINICIANS)
        for k in range(CLINICIANS):
            ws.cell(ws.max_row, 3 + k).fill = FILL
            score.add(ws.cell(ws.max_row, 3 + k))
    last = ws.max_row
    # SUS score: odd items score - 1, even items 5 - score, sum x 2.5 (0-100); blank until all ten are answered.
    row = ["", "SUS score (0-100)"]
    for k in range(CLINICIANS):
        col = chr(ord("C") + k)
        odd = "+".join(f"({col}{r}-1)" for r in range(first, last + 1, 2))
        even = "+".join(f"(5-{col}{r})" for r in range(first + 1, last + 1, 2))
        row.append(f'=IF(COUNT({col}{first}:{col}{last})=10,({odd}+{even})*2.5,"")')
    ws.append(row)

    for r in (1, sus_title):
        ws.cell(r, 1).font = TITLE
    for r in (tasks_header, sus_header, ws.max_row):
        for cell in ws[r]:
            cell.font = HEADER
    for row_cells in ws.iter_rows(min_row=3):
        row_cells[1].alignment = WRAP
    _widths(ws, [5, 70] + [18, 30] * CLINICIANS)
    return len(TASKS) + len(SUS)


def build_pack(cfg: dict, out: Path) -> dict[str, int]:
    """Write the meeting workbook to `out`; returns the number of rows to fill per sheet."""
    wb = Workbook()
    counts = {"0_Read_me": _read_me(wb.active, cfg)}
    wb.active.title = "0_Read_me"
    counts["1_AHP"] = _ahp(wb.create_sheet("1_AHP"))
    counts["2_Thresholds"] = _thresholds(wb.create_sheet("2_Thresholds"), cfg)
    counts["3_Usability"] = _usability(wb.create_sheet("3_Usability"))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return counts
